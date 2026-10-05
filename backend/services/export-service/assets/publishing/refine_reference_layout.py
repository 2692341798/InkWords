"""Update the existing local reference document's navigation and page styles.

Usage: python refine_reference_layout.py INPUT.docx OUTPUT.docx
No downloads; retains the book's existing typefaces, page geometry and styles.
"""
import argparse
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input")
    parser.add_argument("output")
    args = parser.parse_args()
    doc = Document(args.input)
    title = doc.styles["Title"]
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.font.underline = False
    for border in title.element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
    # Keep each short source paragraph together, avoiding isolated note numbers.
    note = doc.styles["Footnote Text"]
    note.font.size = Pt(9)
    note.paragraph_format.keep_together = True
    note.paragraph_format.space_after = Pt(3)
    doc.styles["Source Code"].paragraph_format.widow_control = True
    doc.styles["Source Code"].paragraph_format.keep_together = True
    # WPS otherwise splits a short comparison row between pages even when
    # each cell's paragraph is kept together. Apply the rule at row scope.
    table_style = doc.styles["Table"].element
    row_properties = table_style.find(qn("w:trPr"))
    if row_properties is None:
        row_properties = OxmlElement("w:trPr")
        table_properties = table_style.find(qn("w:tblPr"))
        table_style.insert(table_style.index(table_properties) + 1, row_properties)
    if row_properties.find(qn("w:cantSplit")) is None:
        row_properties.append(OxmlElement("w:cantSplit"))
    for level in (1, 2):
        name = f"TOC {level}"
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.paragraph_format.left_indent = Pt((level - 1) * 14)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_together = True
        style.font.size = Pt(11)
    for section in doc.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run("第 ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        paragraph._p.append(field)
        paragraph.add_run(" 页")
    doc.save(args.output)


if __name__ == "__main__":
    main()
